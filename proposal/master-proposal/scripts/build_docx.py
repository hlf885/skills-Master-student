#!/usr/bin/env python3
"""把开题报告 Markdown 稿转换为排版好的 Word 文档（需要 pandoc 与 python-docx）。

用法:
    python build_docx.py 稿件.md -o 开题报告.docx [--latex-text] [--resource-path 图片目录]

排版（对标参考样本）：
  - 正文：中文宋体 / 西文 Times New Roman，五号（10.5 pt），1.5 倍行距，首行缩进 2 字符；
  - 标题：黑体加粗；图题、表题居中；参考文献条目悬挂缩进；
  - 页边距：上下 2.54 cm、左 2.5 cm、右 1.5 cm；
  - 引用 ^[n]^ 输出为上标；
  - 公式：$$…$$ 转为 Word 公式（OMML），可在 Word 中用 MathType “Convert Equations” 批量转为 MathType 对象；
    公式后单独一行的编号 “(3-1)” 会与公式合并为“居中公式 + 右对齐编号”；
  - --latex-text：公式保留为 $$…$$ 文本，便于在 Word 中用 MathType 的 Toggle TeX 逐个转换。
"""
from __future__ import annotations

import argparse
import copy
import re
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:  # pragma: no cover
    sys.exit("需要 python-docx：pip install python-docx")

BODY_PT = 10.5
CN_FONT, EN_FONT, HEAD_FONT = "宋体", "Times New Roman", "黑体"
EQ_NO = re.compile(r"^\s*[（(](\d+[-.]\d+)[）)]\s*$")
CAPTION = re.compile(r"^\s*(图|表)\s*\d+[-.]\d+")
REF_ITEM = re.compile(r"^\s*\[\d+\]")
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def preprocess(md: str, latex_text: bool) -> str:
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    # 公式中的 \tag{3-1} → 公式后独立一行编号
    md = re.sub(r"\$\$(.*?)\\tag\{([^}]*)\}\s*\$\$", lambda m: f"$${m.group(1).strip()}$$\n\n({m.group(2)})", md, flags=re.S)
    if latex_text:
        md = md.replace("$", r"\$")
    # 参考文献行：“[1] …”在 Markdown 中可能被当成链接引用，前置转义
    md = re.sub(r"^\[(\d+)\]", r"\\[\1\\]", md, flags=re.M)
    # 每条参考文献单独成段
    md = re.sub(r"^(\\\[\d+\\\].*)$", r"\1\n", md, flags=re.M)
    return md


def set_run_font(run, size_pt: float | None = None, east: str = CN_FONT, west: str = EN_FONT, bold: bool | None = None) -> None:
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, fonts)
    for attr, val in (("w:ascii", west), ("w:hAnsi", west), ("w:eastAsia", east), ("w:cs", west)):
        fonts.set(qn(attr), val)
    if size_pt:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def style_fonts(doc) -> None:
    by_name = {st.name: st for st in doc.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Image Caption", "Table Caption", "Caption"):
        st = by_name.get(name)
        if st is None:
            continue
        st.font.name = EN_FONT
        st.font.size = Pt(BODY_PT)
        rpr = st.element.get_or_add_rPr()
        f = rpr.find(qn("w:rFonts"))
        if f is None:
            f = rpr.makeelement(qn("w:rFonts"), {})
            rpr.insert(0, f)
        f.set(qn("w:eastAsia"), CN_FONT)
        f.set(qn("w:ascii"), EN_FONT)
        f.set(qn("w:hAnsi"), EN_FONT)
    for lvl, size in ((1, 12), (2, 10.5), (3, 10.5), (4, 10.5)):
        st = by_name.get(f"Heading {lvl}")
        if st is None:
            continue
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = None
        st.font.name = EN_FONT
        rpr = st.element.get_or_add_rPr()
        f = rpr.find(qn("w:rFonts"))
        if f is None:
            f = rpr.makeelement(qn("w:rFonts"), {})
            rpr.insert(0, f)
        f.set(qn("w:eastAsia"), HEAD_FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if f.get(qn(a)) is not None:
                del f.attrib[qn(a)]
        st.paragraph_format.space_before = Pt(6)
        st.paragraph_format.space_after = Pt(6)


def merge_equation_numbers(doc, text_width_cm: float) -> int:
    paras = doc.paragraphs
    merged = 0
    for i, p in enumerate(paras[:-1]):
        omp = p._element.find(f".//{{{M_NS}}}oMathPara")
        nxt = paras[i + 1]
        if omp is None or not EQ_NO.match(nxt.text):
            continue
        number = "(" + EQ_NO.match(nxt.text).group(1) + ")"
        maths = [copy.deepcopy(m) for m in omp.findall(f"{{{M_NS}}}oMath")]
        omp.getparent().remove(omp)
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.tab_stops.add_tab_stop(Cm(text_width_cm / 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(text_width_cm), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t")
        for m in maths:
            p._element.append(m)
        tail = p.add_run("\t" + number)
        set_run_font(tail, BODY_PT)
        nxt._element.getparent().remove(nxt._element)
        merged += 1
    return merged


def format_document(path: Path) -> None:
    doc = Document(str(path))
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
    text_width = 21 - 2.5 - 1.5
    style_fonts(doc)
    merged = merge_equation_numbers(doc, text_width)

    for p in doc.paragraphs:
        sname = p.style.name if p.style is not None else ""
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        if sname.startswith("Heading"):
            for r in p.runs:
                set_run_font(r, None, east=HEAD_FONT, bold=True)
            continue
        has_math = p._element.find(f".//{{{M_NS}}}oMath") is not None
        has_img = bool(p._element.findall(".//" + qn("w:drawing")))
        if has_img or sname in ("Image Caption", "Table Caption", "Caption") or CAPTION.match(p.text):
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Cm(0)
        elif REF_ITEM.match(p.text):
            pf.first_line_indent = Cm(-0.74)
            pf.left_indent = Cm(0.74)
        elif has_math and (not p.text.strip() or re.search(r"\(\d+[-.]\d+\)\s*$", p.text)):
            pf.first_line_indent = Cm(0)
        else:
            pf.first_line_indent = Pt(BODY_PT * 2)
        for r in p.runs:
            set_run_font(r, BODY_PT)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        set_run_font(r, BODY_PT)
    doc.save(str(path))
    print(f"已合并公式编号 {merged} 处")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown", type=Path)
    ap.add_argument("-o", "--output", type=Path, required=True)
    ap.add_argument("--latex-text", action="store_true", help="公式保留为 LaTeX 文本（供 MathType Toggle TeX 转换）")
    ap.add_argument("--resource-path", default=None, help="图片所在目录（默认与稿件同目录）")
    args = ap.parse_args()

    md = preprocess(args.markdown.read_text(encoding="utf-8"), args.latex_text)
    res = args.resource_path or str(args.markdown.parent.resolve())
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "src.md"
        src.write_text(md, encoding="utf-8")
        cmd = ["pandoc", str(src), "-f", "markdown-inline_notes+superscript",
               "-t", "docx", "-o", str(args.output), "--resource-path", res]
        r = subprocess.run(cmd, capture_output=True, text=True)
        if r.returncode != 0:
            sys.stderr.write(r.stderr)
            return r.returncode
        if r.stderr.strip():
            sys.stderr.write(r.stderr)
    format_document(args.output)
    print(f"已生成 {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
